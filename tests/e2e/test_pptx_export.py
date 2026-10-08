"""PPTX builder: export all 31 sample layouts through the real browser builder and inspect the file.

Needs the built frontend (cd frontend && npm run build) and Playwright Chromium; skipped otherwise.
Starts its own API process that also serves frontend/dist.
"""

import json
import os
import socket
import subprocess
import sys
import time
import urllib.request

import pytest
from conftest import BACKEND, FRONTEND

pptx = pytest.importorskip("pptx")
playwright = pytest.importorskip("playwright.sync_api")
DIST = FRONTEND / "dist"
pytestmark = pytest.mark.skipif(not (DIST / "index.html").exists(), reason="frontend not built")


def _free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


@pytest.fixture(scope="module")
def server(tmp_path_factory):
    port = _free_port()
    env = {**os.environ, "FRONTEND_DIST_DIR": str(DIST), "DATA_DIR": str(tmp_path_factory.mktemp("data"))}
    proc = subprocess.Popen(
        [sys.executable, "-m", "uvicorn", "app.main:app", "--port", str(port), "--log-level", "warning"],
        cwd=BACKEND, env=env,
    )
    url = f"http://127.0.0.1:{port}"
    for _ in range(150):
        try:
            urllib.request.urlopen(url + "/health", timeout=1)
            break
        except OSError:
            time.sleep(0.2)
    yield url
    proc.terminate()
    proc.wait(10)


@pytest.mark.parametrize("theme", ["midnight", "swiss"])
def test_pptx_has_editable_text_native_charts_and_tables(server, theme, tmp_path):
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        try:
            browser = p.chromium.launch()
        except Exception as e:  # chromium not installed
            pytest.skip(f"chromium unavailable: {e}")
        page = browser.new_page(viewport={"width": 1280, "height": 720}, accept_downloads=True)
        page.goto(f"{server}/gallery?native=1&theme={theme}&photos=1", wait_until="networkidle")
        page.evaluate("document.fonts.ready")
        page.wait_for_timeout(500)
        with page.expect_download(timeout=180000) as dl:
            page.click("[data-testid=export-sample-pptx]")
        path = tmp_path / f"{theme}.pptx"
        dl.value.save_as(path)
        browser.close()

    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation(path)
    assert round(Emu(prs.slide_width).inches, 2) == 13.33 and round(Emu(prs.slide_height).inches, 2) == 7.5
    slides = list(prs.slides)
    assert len(slides) == 31

    layouts = json.loads((BACKEND / "app" / "spec" / "layouts.json").read_text(encoding="utf-8"))["layouts"]
    by_layout = dict(zip(layouts, slides))

    for lid, s in by_layout.items():
        assert s.background.fill.type is not None, f"{lid}: no background image"
        assert s.notes_slide.notes_text_frame.text == f"Notes for {lid}"
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        assert texts, f"{lid}: no editable text"

    # Titles survive as editable text (soft line breaks become vertical tabs in python-pptx)
    def all_text(slide):
        return " ".join(sh.text_frame.text.replace("\x0b", " ") for sh in slide.shapes if sh.has_text_frame)

    assert "Pulse: the AI coach that fits your life" in all_text(by_layout["cover-hero"]).replace("PULSE: THE AI COACH THAT FITS YOUR LIFE", "Pulse: the AI coach that fits your life")
    # Native charts and tables
    assert any(sh.has_chart for sh in by_layout["chart-full"].shapes)
    assert any(sh.has_chart for sh in by_layout["chart-insight"].shapes)
    tables = [sh for sh in by_layout["table"].shapes if sh.has_table]
    assert tables and len(tables[0].table.rows) == 6 and len(tables[0].table.columns) == 4
    # Pictures for icons / photos
    assert any(sh.shape_type == 13 for sh in by_layout["icon-grid-3"].shapes)
    assert any(sh.shape_type == 13 for sh in by_layout["cover-split"].shapes)
    # Theme fonts are referenced by name
    fonts = {r.font.name for sh in by_layout["icon-grid-3"].shapes if sh.has_text_frame for para in sh.text_frame.paragraphs for r in para.runs}
    assert fonts & {"Sora", "Inter"}
